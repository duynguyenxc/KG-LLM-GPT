"""Generate a human coding workbook for the per-paper CMOC gold standard.

The per-paper gold standard MUST be produced by independent human coders — not by
the model, and not by AI-adjudicating the model's own output (the circularity that
made the old ``Richmond_Per_Paper_Adjudication_Completed.xlsx`` unusable as ground
truth). This module builds a clean, self-documenting workbook designed for a
defensible, reviewer-proof coding protocol:

  * Sheet ``Instructions`` — the coding protocol and the two-coder + kappa workflow.
  * Sheet ``E-code reference`` — Richmond's 47 entities (the controlled vocabulary).
  * Sheet ``R-code reference`` — Richmond's 40 relation predicates.
  * One sheet per study ``S0xx`` — the coder reads the paper and records CMOCs
    BLIND (model output is deliberately NOT shown, to avoid anchoring bias).
  * Sheet ``Model candidates`` — the model's extractions, provided ONLY for the
    post-blind ADJUDICATION pass, clearly separated.

Cohen's kappa between two coders' completed workbooks is computed by
:func:`compute_kappa`.
"""

from __future__ import annotations

import json
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

from res_pipeline.core.config import GOLD_DIR, OUTPUTS_DIR
from res_pipeline.core.db import get_connection

_HEADER_FILL = PatternFill("solid", fgColor="2F5BD0")
_HEADER_FONT = Font(color="FFFFFF", bold=True)
_WRAP = Alignment(wrap_text=True, vertical="top")


def _style_header(ws, row: int = 1) -> None:
    for cell in ws[row]:
        cell.fill = _HEADER_FILL
        cell.font = _HEADER_FONT
        cell.alignment = Alignment(vertical="center")


def build_coding_workbook(out_path: Path | None = None, coder_label: str = "coder_A") -> Path:
    gold = json.loads((GOLD_DIR / "richmond_gold.json").read_text(encoding="utf-8"))
    out_path = out_path or (OUTPUTS_DIR / "gold" / f"per_paper_coding_{coder_label}.xlsx")
    out_path.parent.mkdir(parents=True, exist_ok=True)

    wb = Workbook()

    # --- Instructions ---
    ws = wb.active
    ws.title = "Instructions"
    ws.column_dimensions["A"].width = 110
    instructions = [
        ("Per-paper CMOC gold-standard coding protocol", True),
        ("", False),
        ("Purpose: build an INDEPENDENT human gold standard of the Context-Mechanism-Outcome "
         "configurations (CMOCs) each of the 28 Richmond studies supports, so the automated "
         "pipeline's per-paper extraction can be validated without circularity.", False),
        ("", False),
        ("Protocol (RAMESES-aligned):", True),
        ("1. Two coders complete a copy of this workbook INDEPENDENTLY (coder_A, coder_B).", False),
        ("2. For each study sheet, READ THE PAPER and record every CMOC you judge it to support. "
         "Do this BLIND — do not look at the 'Model candidates' sheet until step 4.", False),
        ("3. For each CMOC element, pick the closest Richmond E-code from the dropdown (or write "
         "'NEW' + a short label if genuinely outside the E01-E47 vocabulary), and paste the "
         "verbatim supporting quote with its page.", False),
        ("4. ADJUDICATION: after both coders finish blind, compare; compute Cohen's kappa "
         "(res_pipeline.verification.gold_template.compute_kappa). Only THEN consult the "
         "'Model candidates' sheet to check for anything both humans missed; disagreements are "
         "resolved by discussion and logged.", False),
        ("5. The reconciled workbook becomes gold/per_paper_gold/ — the ground truth for the "
         "CMOC-extraction verification metric.", False),
        ("", False),
        ("Element types: Context, Intervention, Mechanism_Resource, Mechanism_Response, Outcome "
         "(see 'E-code reference'). Every CMOC needs at least one Context and one Outcome.", False),
    ]
    for i, (text, bold) in enumerate(instructions, 1):
        cell = ws.cell(row=i, column=1, value=text)
        cell.alignment = _WRAP
        if bold:
            cell.font = Font(bold=True, size=12 if i == 1 else 11)

    # --- E-code reference ---
    ws = wb.create_sheet("E-code reference")
    ws.append(["E-code", "Category", "Label (Richmond)", "Location"])
    _style_header(ws)
    for code, spec in gold["entities"].items():
        ws.append([code, spec["category"], spec["label"], spec.get("location", "")])
    for col, width in zip("ABCD", (10, 22, 80, 28)):
        ws.column_dimensions[col].width = width

    ecode_list = ",".join(list(gold["entities"].keys()) + ["NEW"])

    # --- R-code reference ---
    ws = wb.create_sheet("R-code reference")
    ws.append(["R-code", "Subject", "Predicate", "Object"])
    _style_header(ws)
    cats = {c: s["category"] for c, s in gold["entities"].items()}
    for rel in gold["relationships"]:
        ws.append([rel["id"], f"{rel['subject_code']} ({cats[rel['subject_code']]})",
                   rel["predicate"], f"{rel['object_code']} ({cats[rel['object_code']]})"])
    for col, width in zip("ABCD", (10, 40, 16, 40)):
        ws.column_dimensions[col].width = width

    # --- One coding sheet per study ---
    with get_connection() as conn:
        studies = conn.execute(
            "SELECT study_id, title, year FROM studies ORDER BY study_id"
        ).fetchall()

    for study in studies:
        ws = wb.create_sheet(study["study_id"])
        ws.cell(row=1, column=1,
                value=f"{study['study_id']} ({study['year']}): {study['title']}").font = Font(bold=True)
        ws.append([])
        header = ["CMOC #", "Element type", "E-code", "Your label", "Verbatim quote", "Page", "Polarity"]
        ws.append(header)
        _style_header(ws, row=3)
        dv_ecode = DataValidation(type="list", formula1=f'"{ecode_list[:250]}"', allow_blank=True)
        dv_type = DataValidation(
            type="list",
            formula1='"Context,Intervention,Mechanism_Resource,Mechanism_Response,Outcome"',
            allow_blank=True)
        dv_pol = DataValidation(type="list", formula1='"positive,negative,mixed"', allow_blank=True)
        ws.add_data_validation(dv_ecode)
        ws.add_data_validation(dv_type)
        ws.add_data_validation(dv_pol)
        for r in range(4, 40):
            dv_type.add(ws.cell(row=r, column=2))
            dv_ecode.add(ws.cell(row=r, column=3))
            dv_pol.add(ws.cell(row=r, column=7))
        for col, width in zip("ABCDEFG", (8, 20, 10, 32, 60, 8, 10)):
            ws.column_dimensions[col].width = width

    # --- Model candidates (for adjudication only) ---
    ws = wb.create_sheet("Model candidates")
    ws.append(["DO NOT CONSULT UNTIL BLIND CODING IS COMPLETE (step 4)."])
    ws["A1"].font = Font(bold=True, color="B23B3B")
    ws.append(["Study", "CMOC", "Element type", "Model label", "Verbatim quote", "Support"])
    _style_header(ws, row=2)
    with get_connection() as conn:
        rows = conn.execute(
            "SELECT e.study_id, e.cmoc_id, e.entity_type, e.label, e.verbatim_quote, "
            "c.verifier_support FROM entity_instances e JOIN cmocs c USING (cmoc_id) "
            "ORDER BY e.study_id, e.cmoc_id"
        ).fetchall()
    for row in rows:
        ws.append([row["study_id"], row["cmoc_id"], row["entity_type"], row["label"],
                   (row["verbatim_quote"] or "")[:200], round(row["verifier_support"] or 0, 2)])
    for col, width in zip("ABCDEF", (8, 20, 20, 32, 60, 10)):
        ws.column_dimensions[col].width = width

    wb.save(out_path)
    return out_path


def compute_kappa(coder_a_codes: list[str], coder_b_codes: list[str]) -> float:
    """Cohen's kappa for two aligned lists of categorical codes."""
    if len(coder_a_codes) != len(coder_b_codes) or not coder_a_codes:
        raise ValueError("Coder code lists must be non-empty and equal length.")
    n = len(coder_a_codes)
    observed = sum(a == b for a, b in zip(coder_a_codes, coder_b_codes)) / n
    labels = set(coder_a_codes) | set(coder_b_codes)
    expected = sum(
        (coder_a_codes.count(k) / n) * (coder_b_codes.count(k) / n) for k in labels
    )
    return (observed - expected) / (1 - expected) if expected != 1 else 1.0
